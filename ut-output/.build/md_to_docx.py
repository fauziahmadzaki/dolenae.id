"""Konversi markdown laporan UT ke .docx (parser md minimal, tanpa dependensi eksternal selain python-docx).

Dipakai sekali untuk menghasilkan 01_laporan_UT.docx.
Tidak menyentuh source code aplikasi.
"""
import re
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r"D:\Vs Code\Project\PBL\ut-output\01_laporan_UT.md"
OUT = r"D:\Vs Code\Project\PBL\ut-output\01_laporan_UT.docx"

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def add_para(text, style_name=None):
    p = doc.add_paragraph(style=style_name)
    add_runs(p, text)
    return p

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*.+?\*)")

def add_runs(p, text):
    for token in INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            r = p.add_run(token[2:-2]); r.bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            r = p.add_run(token[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            r = p.add_run(token[1:-1]); r.italic = True
        else:
            p.add_run(token)

def strip_emoji(t):
    return t.replace("🟢", "[KODE]").replace("🔴", "[DATA LAPANGAN]").replace("🟡", "[USULAN]")

lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    stripped = line.strip()

    # Tabel
    if stripped.startswith("|") and i + 1 < len(lines) and set(lines[i+1].strip()) <= set("|-: "):
        header = [c.strip() for c in stripped.strip("|").split("|")]
        i += 2
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        table = doc.add_table(rows=1, cols=len(header))
        table.style = "Light Grid Accent 1"
        for j, h in enumerate(header):
            cell = table.rows[0].cells[j]
            cell.text = ""
            add_runs(cell.paragraphs[0], strip_emoji(h))
            for r in cell.paragraphs[0].runs:
                r.bold = True
        for row in rows:
            cells = table.add_row().cells
            for j in range(len(header)):
                val = strip_emoji(row[j]) if j < len(row) else ""
                cells[j].text = ""
                add_runs(cells[j].paragraphs[0], val)
        doc.add_paragraph()
        continue

    if stripped.startswith("```"):
        i += 1
        buf = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            buf.append(lines[i])
            i += 1
        i += 1
        p = doc.add_paragraph()
        r = p.add_run("\n".join(buf))
        r.font.name = "Consolas"; r.font.size = Pt(9)
        continue

    if not stripped:
        i += 1
        continue

    # Heading
    m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
    if m:
        level = len(m.group(1))
        text = strip_emoji(m.group(2))
        if level == 1:
            h = doc.add_heading(text, level=0)
        else:
            h = doc.add_heading(text, level=min(level, 4))
        i += 1
        continue

    if stripped == "---":
        doc.add_paragraph("─" * 40)
        i += 1
        continue

    # List
    m = re.match(r"^[-*]\s+(.*)$", stripped)
    if m:
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, strip_emoji(m.group(1)))
        i += 1
        continue
    m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
    if m:
        p = doc.add_paragraph(style="List Number")
        add_runs(p, strip_emoji(m.group(2)))
        i += 1
        continue
    if stripped.startswith("> "):
        p = doc.add_paragraph()
        r = p.add_run(strip_emoji(stripped[2:]))
        r.italic = True
        i += 1
        continue

    add_para(strip_emoji(stripped))
    i += 1

doc.save(OUT)
print("OK ->", OUT)
